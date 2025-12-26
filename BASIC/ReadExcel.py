from openpyxl import Workbook,load_workbook

wb = load_workbook(filename="sample_excel.xlsx")

# sh = wb.active(sheet="SampleSheet")
sh = wb["DemoSheet"]
print(sh["A1"].value)
print(sh.cell(row=2,column=2).value)

